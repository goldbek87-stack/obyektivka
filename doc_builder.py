# -*- coding: utf-8 -*-
"""
Yig'ilgan ma'lumotlardan (lat/cyr) tayyor .docx ob'ektivka hujjatini yasaydi.
"""
import os
import uuid

from docx import Document
from docx.shared import Pt, Cm, Mm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

from texts import (
    TITLE, SIMPLE_STEPS, MEHNAT_HEADING, RELATIVES_HEADING, TABLE_HEADERS,
    RELATIVE_PLAN, QUICK_NO_TEXT, L, C,
)
from config import OUTPUT_DIR

FONT_NAME = "Times New Roman"
FONT_SIZE = Pt(12)


def _set_font(run, size=FONT_SIZE, bold=False):
    run.font.name = FONT_NAME
    run.font.size = size
    run.bold = bold
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), FONT_NAME)


def _add_paragraph(doc, text="", align=None, bold=False, size=FONT_SIZE, space_after=4):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    run = p.add_run(text)
    _set_font(run, size=size, bold=bold)
    return p


def _add_label_value(doc, label, value):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    r1 = p.add_run(f"{label}: ")
    _set_font(r1, bold=True)
    r2 = p.add_run(value if value else "—")
    _set_font(r2, bold=False)
    return p


def _set_cell_borders_none(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "nil")
        borders.append(el)
    tcPr.append(borders)


def _style_table_borders(table):
    tbl = table._tbl
    tblPr = tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "000000")
        borders.append(el)
    tblPr.append(borders)


def build_document(lang: str, data: dict, mehnat: list, relatives: list,
                    photo_path: str | None) -> str:
    """
    lang: 'lat' yoki 'cyr'
    data: SIMPLE_STEPS kalitlariga mos oddiy maydonlar (fio bundan mustasno -
          sarlavha sifatida ishlatiladi)
    mehnat: [{'from':..,'to':..,'place':..}, ...]
    relatives: [{'col':..,'fio':..,'tug':..,'ish':..,'turar':..}, ...]
    photo_path: rasmga yo'l (bo'lmasa None)

    Qaytaradi: tayyor .docx faylining to'liq yo'li
    """
    doc = Document()

    section = doc.sections[0]
    section.top_margin = Cm(1.5)
    section.bottom_margin = Cm(1.5)
    section.left_margin = Cm(2)
    section.right_margin = Cm(1.5)

    fio = data.get("fio", "").strip()

    # --- Sarlavha + rasm (yuqori jadval: chap - sarlavha/ism, o'ng - rasm) ---
    header_table = doc.add_table(rows=1, cols=2)
    header_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    header_table.autofit = False
    left_cell, right_cell = header_table.rows[0].cells
    _set_cell_borders_none(left_cell)
    _set_cell_borders_none(right_cell)
    left_cell.width = Cm(13)
    right_cell.width = Cm(3.5)

    left_cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    lp = left_cell.paragraphs[0]
    lp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = lp.add_run(TITLE[lang])
    _set_font(r, size=Pt(14), bold=True)
    lp2 = left_cell.add_paragraph()
    lp2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = lp2.add_run(fio)
    _set_font(r2, size=Pt(13), bold=True)

    right_cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    rp = right_cell.paragraphs[0]
    rp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if photo_path and os.path.exists(photo_path):
        run = rp.add_run()
        run.add_picture(photo_path, width=Mm(30), height=Mm(40))
    else:
        r3 = rp.add_run("[ расм ]" if lang == C else "[ rasm ]")
        _set_font(r3)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # --- Oddiy maydonlar ---
    for step in SIMPLE_STEPS:
        if step.get("is_header"):
            continue
        key = step["key"]
        label = step["label"][0] if lang == L else step["label"][1]
        value = data.get(key, "")
        _add_label_value(doc, label, value)

    # --- Mehnat faoliyati ---
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    _add_paragraph(doc, MEHNAT_HEADING[lang], bold=True, size=Pt(13), space_after=6)
    if not mehnat:
        _add_paragraph(doc, QUICK_NO_TEXT[lang].capitalize(), space_after=3)
    yy = "yy." if lang == L else "йй."
    for entry in mehnat:
        frm = entry.get("from", "").strip()
        to = entry.get("to", "").strip()
        place = entry.get("place", "").strip()
        line = f"{frm}-{to} {yy}  {place}"
        _add_paragraph(doc, line, space_after=3)

    # --- Yaqin qarindoshlar ---
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    heading_lines = (f"{fio}ning " if lang == L else f"{fio}нинг ") + RELATIVES_HEADING[lang]
    for line in heading_lines.split("\n"):
        _add_paragraph(doc, line, align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=Pt(13), space_after=2)

    headers = TABLE_HEADERS[lang]
    rel_table = doc.add_table(rows=1, cols=len(headers))
    rel_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    _style_table_borders(rel_table)
    hdr_cells = rel_table.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = ""
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h)
        _set_font(r, size=Pt(11), bold=True)

    for rel in relatives:
        row_cells = rel_table.add_row().cells
        values = [
            rel.get("col", ""),
            rel.get("fio", ""),
            rel.get("tug", ""),
            rel.get("ish", ""),
            rel.get("turar", ""),
        ]
        for i, v in enumerate(values):
            row_cells[i].text = ""
            p = row_cells[i].paragraphs[0]
            r = p.add_run(v)
            _set_font(r, size=Pt(11))

    out_name = f"{uuid.uuid4().hex}.docx"
    out_path = os.path.join(OUTPUT_DIR, out_name)
    doc.save(out_path)
    return out_path
